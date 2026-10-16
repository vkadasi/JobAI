"""Read uploaded resumes and render structured resumes to DOCX and PDF."""

import io
import os
import re
from xml.sax.saxutils import escape

# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------


def extract_text(filename, data):
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(lines)
    elif ext in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Please upload a PDF, DOCX or TXT file.")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 50:
        raise ValueError("Couldn't read text from this file (it may be a scanned image). Try a DOCX or paste the text instead.")
    return text


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _date_range(item):
    start, end = item.get("start", "").strip(), item.get("end", "").strip()
    if start and end:
        return f"{start} – {end}"
    return start or end


def _contact_line(resume):
    c = resume.get("contact", {})
    return "  |  ".join(v.strip() for v in (c.get("email"), c.get("phone"), c.get("location"), c.get("linkedin"), c.get("website")) if v and v.strip())


def _join(*parts, sep=", "):
    return sep.join(p.strip() for p in parts if p and p.strip())


def safe_filename(resume, analysis=None):
    name = resume.get("name") or "Resume"
    company = (analysis or {}).get("company", "")
    base = "_".join(x for x in [name, "Resume", company] if x)
    return re.sub(r"[^A-Za-z0-9_-]+", "_", base).strip("_") or "Resume"


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def build_docx(resume, paper="A4"):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = docx.Document()
    section = doc.sections[0]
    if paper == "Letter":
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    else:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.7))
    section.top_margin = section.bottom_margin = Inches(0.6)
    content_width = section.page_width - section.left_margin - section.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.05

    accent = RGBColor(0x1F, 0x3A, 0x5F)

    def para(text="", bold=False, size=None, align=None, space_before=0, space_after=0, italic=False, color=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if align:
            p.alignment = align
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def heading(title):
        p = para(title.upper(), bold=True, size=11, space_before=9, space_after=3, color=accent)
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(k), v)
        bdr.append(bottom)
        pPr.append(bdr)

    def left_right(left_bold, left_rest, right, space_before=4):
        p = para(space_before=space_before)
        p.paragraph_format.tab_stops.add_tab_stop(content_width, WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(left_bold)
        r.bold = True
        if left_rest:
            p.add_run(left_rest)
        if right:
            p.add_run("\t" + right)
        return p

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.25)
        p.add_run(text)

    # Header
    para(resume.get("name", ""), bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, color=accent)
    if resume.get("headline"):
        para(resume["headline"], size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1)
    contact = _contact_line(resume)
    if contact:
        para(contact, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2)

    if resume.get("summary"):
        heading("Summary")
        para(resume["summary"])

    if resume.get("skills"):
        heading("Skills")
        for s in resume["skills"]:
            p = para(space_after=1)
            if s.get("category"):
                p.add_run(s["category"] + ": ").bold = True
            p.add_run(", ".join(s.get("items", [])))

    if resume.get("experience"):
        heading("Experience")
        for i, e in enumerate(resume["experience"]):
            left_right(e.get("title", ""), (" — " + e["company"]) if e.get("company") else "", _date_range(e), space_before=0 if i == 0 else 6)
            if e.get("location"):
                para(e["location"], italic=True, size=9.5)
            for b in e.get("bullets", []):
                bullet(b)

    if resume.get("projects"):
        heading("Projects")
        for i, pr in enumerate(resume["projects"]):
            left_right(pr.get("name", ""), "", pr.get("link", ""), space_before=0 if i == 0 else 5)
            for b in pr.get("bullets", []):
                bullet(b)

    if resume.get("education"):
        heading("Education")
        for ed in resume["education"]:
            left_right(ed.get("degree", ""), (" — " + ed["school"]) if ed.get("school") else "", _date_range(ed), space_before=2)
            extra = _join(ed.get("location", ""), ed.get("details", ""), sep="  |  ")
            if extra:
                para(extra, size=9.5)

    if resume.get("certifications"):
        heading("Certifications")
        for c in resume["certifications"]:
            bullet(c)

    for sec in resume.get("extra_sections", []):
        if sec.get("title") and sec.get("items"):
            heading(sec["title"])
            for it in sec["items"]:
                bullet(it)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_FONT_CANDIDATES = [
    # Windows
    ("C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/calibrib.ttf", "C:/Windows/Fonts/calibrii.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
    # macOS
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/System/Library/Fonts/Supplemental/Arial Italic.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf", "/Library/Fonts/Arial Italic.ttf"),
    # Linux
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
]
_fonts = None


def _register_fonts():
    """Use a system TrueType font (for full Unicode); fall back to built-in Helvetica."""
    global _fonts
    if _fonts:
        return _fonts
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold, italic in _FONT_CANDIDATES:
        if all(os.path.exists(p) for p in (regular, bold, italic)):
            try:
                pdfmetrics.registerFont(TTFont("ResumeFont", regular))
                pdfmetrics.registerFont(TTFont("ResumeFont-Bold", bold))
                pdfmetrics.registerFont(TTFont("ResumeFont-Italic", italic))
                addMapping("ResumeFont", 0, 0, "ResumeFont")
                addMapping("ResumeFont", 1, 0, "ResumeFont-Bold")
                addMapping("ResumeFont", 0, 1, "ResumeFont-Italic")
                addMapping("ResumeFont", 1, 1, "ResumeFont-Bold")
                _fonts = ("ResumeFont", "ResumeFont-Bold", "ResumeFont-Italic", True)
                return _fonts
            except Exception:
                continue
    _fonts = ("Helvetica", "Helvetica-Bold", "Helvetica-Oblique", False)
    return _fonts


def build_pdf(resume, paper="A4"):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold, italic, unicode_ok = _register_fonts()
    accent = colors.HexColor("#1F3A5F")

    def t(text):
        text = text or ""
        if not unicode_ok:
            text = text.translate(str.maketrans({"–": "-", "—": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "•": "-"}))
            text = text.encode("latin-1", "replace").decode("latin-1")
        return escape(text)

    base = ParagraphStyle("base", fontName=regular, fontSize=10, leading=13)
    name_s = ParagraphStyle("name", parent=base, fontName=bold, fontSize=20, leading=24, alignment=TA_CENTER, textColor=accent)
    headline_s = ParagraphStyle("headline", parent=base, fontSize=11.5, leading=15, alignment=TA_CENTER)
    contact_s = ParagraphStyle("contact", parent=base, fontSize=9, leading=12, alignment=TA_CENTER, textColor=colors.HexColor("#333333"))
    head_s = ParagraphStyle("head", parent=base, fontName=bold, fontSize=11, leading=14, textColor=accent, spaceBefore=9)
    small_s = ParagraphStyle("small", parent=base, fontSize=9, leading=12, textColor=colors.HexColor("#444444"))
    right_s = ParagraphStyle("right", parent=base, alignment=2)
    bullet_s = ParagraphStyle("bullet", parent=base, leading=12.5)

    pagesize = letter if paper == "Letter" else A4
    width = pagesize[0] - 1.4 * inch
    story = []

    def heading(title):
        story.append(Paragraph(t(title.upper()), head_s))
        story.append(HRFlowable(width="100%", thickness=0.7, color=accent, spaceBefore=1, spaceAfter=4))

    def row(left_html, right_text, space_before=4):
        tbl = Table([[Paragraph(left_html, base), Paragraph(t(right_text), right_s)]], colWidths=[width * 0.72, width * 0.28])
        tbl.setStyle(TableStyle([
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), space_before), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(tbl)

    def bullets(items):
        if not items:
            return
        story.append(ListFlowable(
            [ListItem(Paragraph(t(b), bullet_s), leftIndent=12, value="•" if unicode_ok else "-") for b in items],
            bulletType="bullet", start="•" if unicode_ok else "-", leftIndent=12, bulletFontName=regular, bulletFontSize=9,
            spaceBefore=1,
        ))

    story.append(Paragraph(t(resume.get("name", "")), name_s))
    if resume.get("headline"):
        story.append(Paragraph(t(resume["headline"]), headline_s))
    contact = _contact_line(resume)
    if contact:
        story.append(Spacer(1, 2))
        story.append(Paragraph(t(contact), contact_s))

    if resume.get("summary"):
        heading("Summary")
        story.append(Paragraph(t(resume["summary"]), base))

    if resume.get("skills"):
        heading("Skills")
        for s in resume["skills"]:
            label = f"<b>{t(s['category'])}:</b> " if s.get("category") else ""
            story.append(Paragraph(label + t(", ".join(s.get("items", []))), ParagraphStyle("sk", parent=base, spaceAfter=1)))

    if resume.get("experience"):
        heading("Experience")
        for i, e in enumerate(resume["experience"]):
            left = f"<b>{t(e.get('title', ''))}</b>" + (f" — {t(e['company'])}" if e.get("company") else "")
            row(left, _date_range(e), space_before=0 if i == 0 else 6)
            if e.get("location"):
                story.append(Paragraph(f"<i>{t(e['location'])}</i>", small_s))
            bullets(e.get("bullets", []))

    if resume.get("projects"):
        heading("Projects")
        for i, pr in enumerate(resume["projects"]):
            row(f"<b>{t(pr.get('name', ''))}</b>", pr.get("link", ""), space_before=0 if i == 0 else 5)
            bullets(pr.get("bullets", []))

    if resume.get("education"):
        heading("Education")
        for i, ed in enumerate(resume["education"]):
            left = f"<b>{t(ed.get('degree', ''))}</b>" + (f" — {t(ed['school'])}" if ed.get("school") else "")
            row(left, _date_range(ed), space_before=0 if i == 0 else 3)
            extra = _join(ed.get("location", ""), ed.get("details", ""), sep="  |  ")
            if extra:
                story.append(Paragraph(t(extra), small_s))

    if resume.get("certifications"):
        heading("Certifications")
        bullets(resume["certifications"])

    for sec in resume.get("extra_sections", []):
        if sec.get("title") and sec.get("items"):
            heading(sec["title"])
            bullets(sec["items"])

    buf = io.BytesIO()
    title = resume.get("name", "Resume") + " - Resume"
    SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=0.7 * inch, rightMargin=0.7 * inch,
                      topMargin=0.55 * inch, bottomMargin=0.55 * inch, title=title,
                      author=resume.get("name", "")).build(story)
    return buf.getvalue()
